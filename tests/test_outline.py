"""The PPT stream: authoritative session numbering, and the slide text itself.

The bug this locks down was silent and wrong in the reviewer's hands.
`portal.is_session()` infers session numbers by POSITION - it counts LEARNING_SET units
carrying an INTERACTIVE_VIDEO. Intro to Gen AI has a unit called `Common Mistakes` that
is exactly that shape and is NOT a numbered session, so every session from position 8
onward was reported one too high: 81 of 104 units disagreed with the curriculum's own
numbering, and `expect_sessions` was calibrated to the wrong count, so the integrity
check passed while the digest pointed reviewers at the wrong slide deck.

The workbook is the curriculum team's own record of what a session IS, so where it
speaks it wins. Where it is silent - PSE has no workbook at all - the positional walk
still runs.
"""
import json
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from miw.ingest import outline as O
from miw.ingest.portal import read_course
from miw.ingest.sheets import course_for_workbook

pytest.importorskip("openpyxl")
import openpyxl


# A course whose 2nd unit looks like a session but is not numbered as one - the
# `Common Mistakes` shape, minimised.
COURSE = [{
    "course_title": "Demo",
    "topics": [{"topic_name": "T1", "order": 1, "units": [
        {"unit_id": "u-s1", "unit_name": "Real Session One", "unit_type": "LEARNING_SET",
         "order": 1, "contents": [
             {"learning_resource_type": "INTERACTIVE_VIDEO"},
             {"learning_resource_id": "c1", "content": "uses toolzz here"}]},
        {"unit_id": "u-extra", "unit_name": "Common Mistakes", "unit_type": "LEARNING_SET",
         "order": 2, "contents": [
             {"learning_resource_type": "INTERACTIVE_VIDEO"},
             {"learning_resource_id": "c2", "content": "aside about toolzz"}]},
        {"unit_id": "u-s2", "unit_name": "Real Session Two", "unit_type": "LEARNING_SET",
         "order": 3, "contents": [
             {"learning_resource_type": "INTERACTIVE_VIDEO"},
             {"learning_resource_id": "c3", "content": "more toolzz"}]},
        {"unit_id": "u-quiz", "unit_name": None, "unit_type": "QUIZ", "order": 4,
         "contents": [{"question_id": "q1", "object_type": "OBJECTIVE_QUESTIONS",
                       "content": "quiz about toolzz"}]},
    ]}]}]


def _workbook(path, sessions, practice):
    """Build a workbook with the two PPT-level sheets."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = O.OUTLINE_SHEET
    ws.append(["Course Id", "Course Name", "Topic Name", "Session No.", "Session ID",
               "Session Name", "Key Takeaways", "Outline", "Session PPT",
               "Reading Material id", "Reading Material Content",
               "Recorded Session Transcript"])
    for no, sid, name, takeaways, out in sessions:
        ws.append(["c", "Demo", "T1", no, sid, name, takeaways, out,
                   f"https://slides.example/{sid}", "", "", ""])
    ps = wb.create_sheet(O.PRACTICE_SHEET)
    ps.append(["Topic ID", "Topic Name", "Session ID", "Session Name", "Unit ID",
               "Reading Material/ MCQ / Coding practice"])
    for sid, uid, art in practice:
        ps.append(["t", "T1", sid, "", uid, art])
    wb.save(path)


@pytest.fixture()
def book(tmp_path):
    p = tmp_path / "Demo - Course Contents.xlsx"
    _workbook(p,
              sessions=[(1, "sid-1", "Real Session One", "Learn toolzz", "- toolzz intro"),
                        (2, "sid-2", "Real Session Two", "More toolzz", "- toolzz deep dive")],
              practice=[("sid-1", "u-s1", "Reading Material"),
                        ("sid-2", "u-s2", "Reading Material"),
                        ("sid-2", "u-quiz", "Module Quiz")])
    return p


@pytest.fixture()
def export(tmp_path):
    p = tmp_path / "demo.json"
    p.write_text(json.dumps(COURSE))
    return p


# ------------------------------------------------------------------ the reader

def test_the_outline_sheet_yields_numbered_sessions_with_their_slide_text(book):
    o = O.read_outline(book)
    assert o.stats.errors == []
    assert [s.session_no for s in o.sessions] == [1, 2]
    assert o.sessions[0].session_name == "Real Session One"
    assert "toolzz intro" in o.sessions[0].outline
    assert o.sessions[0].ppt_url.startswith("https://slides.example/")
    assert o.stats.with_outline == 2 and o.stats.with_takeaways == 2


def test_the_practice_sheet_maps_units_to_the_session_they_came_from(book):
    """This is the lineage the export cannot express: which quiz came from which deck."""
    o = O.read_outline(book)
    assert o.session_of_unit == {"u-s1": 1, "u-s2": 2, "u-quiz": 2}
    assert o.stats.unmapped_practice_rows == 0


def test_headers_are_resolved_by_alias_not_by_position(tmp_path):
    """The hand-built and generated workbooks disagree on wording and column order."""
    p = tmp_path / "alias.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = O.OUTLINE_SHEET
    ws.append(["Session Name", "Module Name", "Session No", "Session ID", "Outline"])
    ws.append(["S One", "T1", 1, "sid-1", "- a line"])
    wb.save(p)
    o = O.read_outline(p)
    assert [(s.session_no, s.session_name, s.topic_name) for s in o.sessions] \
        == [(1, "S One", "T1")]


def test_a_workbook_with_no_session_number_column_says_so_and_numbers_nothing(tmp_path):
    p = tmp_path / "bad.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = O.OUTLINE_SHEET
    ws.append(["Session Name", "Outline"])
    ws.append(["S One", "- a line"])
    wb.save(p)
    o = O.read_outline(p)
    assert o.sessions == [] and o.session_of_unit == {}
    assert any("Session No." in e for e in o.stats.errors)


def test_a_missing_file_is_an_error_not_an_exception(tmp_path):
    o = O.read_outline(tmp_path / "nope.xlsx")
    assert o.sessions == [] and o.stats.errors


# ------------------------------------------------- the numbering it corrects

def test_without_the_workbook_the_positional_walk_counts_the_extra_unit(export):
    """The bug, reproduced: `Common Mistakes` becomes session 2 and shifts the rest."""
    recs, st = read_course(str(export), "Demo")
    assert st.sessions == 3, "three units look like a session"
    by_unit = {r.unit_id: r.session_no for r in recs}
    assert by_unit["u-s2"] == 3, "the real second session is reported as the third"
    assert st.session_from_workbook == 0


def test_with_the_workbook_the_curriculums_own_numbering_wins(book, export):
    o = O.read_outline(book)
    recs, st = read_course(str(export), "Demo", o.session_of_unit)
    by_unit = {r.unit_id: r.session_no for r in recs}
    assert by_unit["u-s1"] == 1
    assert by_unit["u-s2"] == 2, "corrected from the positional 3"
    assert by_unit["u-quiz"] == 2, "a quiz belongs to the session it was authored from"
    assert st.session_from_workbook == 3
    assert st.session_conflicts == 2, "u-s2 and u-quiz both disagreed"
    assert st.authoritative_sessions == 2


def test_a_disagreement_is_counted_and_shown_never_silently_resolved(book, export):
    """Silently changing 81 session numbers is how a reviewer stops trusting a digest."""
    o = O.read_outline(book)
    _recs, st = read_course(str(export), "Demo", o.session_of_unit)
    assert st.conflict_examples, "the reason must be reportable"
    assert any("workbook s2" in c and "position s3" in c for c in st.conflict_examples)


def test_units_the_workbook_does_not_cover_still_get_a_positional_number(export, tmp_path):
    """PSE has no workbook at all; a partial workbook must not blank the rest."""
    p = tmp_path / "partial.xlsx"
    _workbook(p, sessions=[(1, "sid-1", "Real Session One", "k", "- o")],
              practice=[("sid-1", "u-s1", "Reading Material")])
    o = O.read_outline(p)
    recs, st = read_course(str(export), "Demo", o.session_of_unit)
    by_unit = {r.unit_id: r.session_no for r in recs}
    assert by_unit["u-s1"] == 1, "from the workbook"
    assert by_unit["u-s2"], "still numbered, by position"
    assert st.session_inferred > 0


# ------------------------------------------------------------ the slide text

def test_the_slide_text_becomes_records_the_extractor_can_read(book):
    """`markdown` because that is what `feed_prose` reads; the outline is authored prose."""
    o = O.read_outline(book)
    recs = O.outline_records(o, "Demo")
    assert len(recs) == 4, "outline + takeaways for each of two sessions"
    assert {r.evidence_source for r in recs} == {"markdown"}
    assert {r.object_type for r in recs} == {"SESSION_PPT"}
    assert [r.session_no for r in recs] == [1, 1, 2, 2]


def test_a_slide_record_points_at_the_workbook_not_a_json_path(book):
    """So the detail panel explains there is no export excerpt instead of rendering blank."""
    o = O.read_outline(book)
    r = O.outline_records(o, "Demo")[0]
    assert "::" in r.field_path and O.OUTLINE_SHEET in r.field_path
    from miw.extract.locate import parse_path
    assert parse_path(r.field_path) is None


def test_an_empty_outline_cell_produces_no_record(tmp_path):
    p = tmp_path / "sparse.xlsx"
    _workbook(p, sessions=[(1, "sid-1", "S One", "", ""),
                           (2, "sid-2", "S Two", "Learn it", "- a real line")],
              practice=[])
    o = O.read_outline(p)
    recs = O.outline_records(o, "Demo")
    assert [r.session_no for r in recs] == [2, 2], "session 1 contributed nothing"
    assert {r.title for r in recs} == {"Outline", "Key Takeaways"}


# ------------------------------------------------------ the shared course matcher

def test_a_workbook_maps_to_a_course_by_normalised_stem():
    assert course_for_workbook("Intro to Generative AI - Course Contents.xlsx") \
        == "Intro to Gen AI"
    assert course_for_workbook("AI for Finance - Course Contents.xlsx") == "AI for Finance"


def test_an_unknown_workbook_maps_to_nothing_rather_than_inventing_a_course():
    """The fallback to the filename is what created three phantom courses."""
    assert course_for_workbook("Some Other Book.xlsx") == ""


# ------------------------------------- sheet declarations get a session number

def test_a_sheet_declaration_is_placed_on_a_session_when_the_name_resolves():
    """Before this, every sheet-derived Location had `session_no = None`.

    43 dependencies in Intro to Gen AI were known only from a tool sheet, so the UI
    could say the course used `Cerebras` but never which session - it rendered
    "workbook" where a session number belonged. The tool sheets name the session in
    their own words, which the ingested records can resolve.
    """
    from miw.extract.inventory import InventoryBuilder
    from miw.ingest.sheets import SheetTool
    from miw.registry import Entry, Registry

    reg = Registry(entries=[Entry(canonical_name="Cerebras", kind="tool",
                                  aliases=["cerebras"])])
    b = InventoryBuilder(reg)
    tools = [SheetTool(name="Cerebras", session="Mastering Image Generation",
                       sheet="Entity Ids - Tools & Versions U", workbook="Demo.xlsx")]
    b.feed_sheets(tools, {"Demo.xlsx": "Demo"},
                  {("Demo", "mastering image generation"): 15})
    dep = next(d for d in b.finish() if d.canonical_name == "Cerebras")
    loc = dep.locations[0]
    assert loc.session_no == 15
    assert loc.evidence_source == "sheet_declared"


def test_an_unresolvable_session_name_leaves_the_number_empty_not_wrong():
    """A guessed session number is worse than none: it sends a reviewer to a deck that
    never mentioned the tool."""
    from miw.extract.inventory import InventoryBuilder
    from miw.ingest.sheets import SheetTool
    from miw.registry import Entry, Registry

    reg = Registry(entries=[Entry(canonical_name="Cerebras", kind="tool")])
    b = InventoryBuilder(reg)
    tools = [SheetTool(name="Cerebras", session="A Session Nobody Ingested",
                       sheet="s", workbook="Demo.xlsx")]
    b.feed_sheets(tools, {"Demo.xlsx": "Demo"}, {("Demo", "something else"): 3})
    dep = next(d for d in b.finish() if d.canonical_name == "Cerebras")
    assert dep.locations[0].session_no is None


# --------------------------------------------- the real workbooks, as shipped

REAL = Path(__file__).resolve().parents[1] / "data" / "sheets"


@pytest.mark.skipif(not REAL.exists(), reason="course workbooks not present")
def test_every_shipped_workbook_still_parses_into_sessions_and_a_unit_map():
    """A workbook renamed or re-shaped upstream must fail loudly here, not silently
    fall back to positional numbering in production."""
    books = sorted(REAL.glob("*.xlsx"))
    assert books, "no workbooks found"
    for b in books:
        o = O.read_outline(b)
        assert o.stats.errors == [], f"{b.name}: {o.stats.errors}"
        assert o.sessions, f"{b.name}: no numbered sessions"
        assert o.session_of_unit, f"{b.name}: no unit -> session mapping"
        assert o.stats.unmapped_practice_rows == 0, f"{b.name}: unmapped practice rows"
        nos = [s.session_no for s in o.sessions]
        assert nos == sorted(nos) and len(set(nos)) == len(nos), \
            f"{b.name}: session numbers are not a clean sequence"


@pytest.mark.skipif(not REAL.exists(), reason="course workbooks not present")
def test_declared_session_counts_match_the_roster():
    """`expect_sessions` must track the CURRICULUM's count, not the export's.

    It was previously calibrated to the export's positional count, which is exactly why
    the off-by-one passed the integrity check.
    """
    from config.constants import COURSES
    for b in sorted(REAL.glob("*.xlsx")):
        title = course_for_workbook(b.name)
        if not title:
            continue
        meta = next((m for m in COURSES.values() if m["title"] == title), None)
        if not meta:
            continue
        assert O.read_outline(b).session_count == meta["expect_sessions"], \
            f"{title}: workbook and config/constants.py disagree"
