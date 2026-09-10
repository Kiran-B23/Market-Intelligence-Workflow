"""The PPT stream: what each session's slide deck covers, and which units come from it.

The JSON export is the *published* course. The slide decks are what the curriculum is
authored FROM — the MCQs, coding questions and reading materials are all written against
a session's deck — and the decks themselves are Google Slides links, so their content
never reaches the export. What does reach us is the workbook's `Course Outline` sheet,
where a curriculum engineer records each session's outline and key takeaways by hand.

Two distinct things are recovered here, and the second is why this module exists at all.

**1. The outline text.** Measured on Intro to Gen AI: of 24 sessions with an outline,
**0** have their outline fully present anywhere in the JSON export, and the `Session PPT`
links appear in it 0 times out of 25. So this is the only textual record of the slides
we have. It names tools at slide level — `Gamma AI`, `Cerebras`, `NVIDIA NIM` — some of
which the export never mentions.

**2. The authoritative session numbering, which the export gets wrong.**
`portal.is_session()` infers a session by *position*: it counts LEARNING_SET units that
carry an INTERACTIVE_VIDEO. That is a good heuristic and it is measurably not the
curriculum's own numbering:

    Intro to Gen AI   portal 26 sessions, workbook 25
                      `Common Mistakes` is a LEARNING_SET with a video at position 8,
                      and is NOT a numbered session -> every session from 8 onward was
                      reported ONE HIGHER than the curriculum's own number.
                      81 of 104 units disagreed.
    AI for Finance    portal 18, workbook 17 (`AI Finance Add-On Session`, last)
    LLM Applications  portal 29, workbook 29 — agrees exactly

That off-by-one reached the reviewer: a digest saying "the earliest affected session is
session 11" pointed at the deck for session 10. The workbook is the curriculum team's
own record of what a session IS, so where it speaks it wins, and `portal.read_course`
takes this mapping as an override rather than guessing.

The chain that makes it possible joins cleanly (104 of 104 rows on Intro to Gen AI):

    Course Outline                  Session ID -> Session No., outline, PPT link
    Session-Practice Content Linked Session ID -> Unit ID + artifact type
    the JSON export                 Unit ID

which is also the *lineage* the export cannot express: it says which MCQ and coding
units were built from which session's deck, even where the unit's own text never names
the tool the deck taught.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from miw.schema import ContentRecord

try:
    import openpyxl
except ImportError:                                        # pragma: no cover
    openpyxl = None

OUTLINE_SHEET = "Course Outline"
PRACTICE_SHEET = "Session-Practice Content Linked"

# Resolved by alias, never by index: the hand-built and generated workbooks disagree on
# header wording and on trailing spaces.
_H = {
    "session_no": ("session no.", "session no", "session number"),
    "session_id": ("session id",),
    "session_name": ("session name", "session title"),
    "topic_name": ("topic name", "module name"),
    "outline": ("outline",),
    "takeaways": ("key takeaways", "key takeaway"),
    "ppt": ("session ppt", "ppt", "session ppt link"),
    "unit_id": ("unit id",),
    "artifact": ("reading material/ mcq / coding practice",
                 "reading material / mcq / coding practice"),
}


def _norm(v) -> str:
    return re.sub(r"\s+", " ", str(v or "").strip().lower())


def _cols(header: list) -> dict:
    """field name -> column index, for whichever aliases this sheet happens to use."""
    norm = [_norm(h) for h in header]
    out = {}
    for field_name, aliases in _H.items():
        for i, h in enumerate(norm):
            if h in aliases:
                out[field_name] = i
                break
    return out


@dataclass
class SessionOutline:
    """One numbered session as the curriculum team records it."""
    session_no: int
    session_id: str = ""
    session_name: str = ""
    topic_name: str = ""
    outline: str = ""
    key_takeaways: str = ""
    ppt_url: str = ""
    row: int = 0


@dataclass
class OutlineStats:
    workbook: str = ""
    sessions: int = 0
    with_outline: int = 0
    with_takeaways: int = 0
    practice_rows: int = 0
    units_mapped: int = 0
    unmapped_practice_rows: int = 0
    errors: list[str] = field(default_factory=list)


@dataclass
class Outline:
    sessions: list[SessionOutline] = field(default_factory=list)
    # unit_id -> the session number that unit was authored from. The override
    # `portal.read_course` consults.
    session_of_unit: dict[str, int] = field(default_factory=dict)
    stats: OutlineStats = field(default_factory=OutlineStats)

    @property
    def session_count(self) -> int:
        return len(self.sessions)


def read_outline(path: str | Path) -> Outline:
    """Read one workbook's PPT-level sheets. Never raises; surprises land in stats."""
    out = Outline(stats=OutlineStats(workbook=Path(path).name))
    if openpyxl is None:                                   # pragma: no cover
        out.stats.errors.append("openpyxl not installed; outline input skipped")
        return out
    try:
        wb = openpyxl.load_workbook(str(path), data_only=True)
    except Exception as exc:
        out.stats.errors.append(f"{Path(path).name}: {type(exc).__name__}: {exc}")
        return out

    # --- the sessions themselves -----------------------------------------
    by_session_id: dict[str, int] = {}
    if OUTLINE_SHEET in wb.sheetnames:
        rows = wb[OUTLINE_SHEET].iter_rows(values_only=True)
        try:
            cols = _cols(list(next(rows)))
        except StopIteration:
            cols = {}
        if "session_no" not in cols:
            out.stats.errors.append(
                f"{OUTLINE_SHEET}: no 'Session No.' column; session numbering cannot be "
                f"taken from this workbook")
        else:
            def cell(row, key):
                i = cols.get(key)
                return "" if i is None or i >= len(row) or row[i] is None else str(row[i]).strip()
            for ri, row in enumerate(rows, start=2):
                raw = row[cols["session_no"]] if cols["session_no"] < len(row) else None
                if raw is None or str(raw).strip() == "":
                    continue
                try:
                    # Excel hands these back as floats.
                    no = int(float(str(raw).strip()))
                except ValueError:
                    out.stats.errors.append(
                        f"{OUTLINE_SHEET} row {ri}: session number {raw!r} is not a number")
                    continue
                s = SessionOutline(
                    session_no=no, session_id=cell(row, "session_id"),
                    session_name=cell(row, "session_name"),
                    topic_name=cell(row, "topic_name"),
                    outline=cell(row, "outline"),
                    key_takeaways=cell(row, "takeaways"),
                    ppt_url=cell(row, "ppt"), row=ri)
                out.sessions.append(s)
                if s.session_id:
                    by_session_id[s.session_id] = no
            out.sessions.sort(key=lambda s: s.session_no)
            out.stats.sessions = len(out.sessions)
            out.stats.with_outline = sum(1 for s in out.sessions if len(s.outline) > 2)
            out.stats.with_takeaways = sum(1 for s in out.sessions
                                           if len(s.key_takeaways) > 2)
    else:
        out.stats.errors.append(f"no '{OUTLINE_SHEET}' sheet in {Path(path).name}")

    # --- session -> the units authored from it ----------------------------
    if PRACTICE_SHEET in wb.sheetnames:
        rows = wb[PRACTICE_SHEET].iter_rows(values_only=True)
        try:
            cols = _cols(list(next(rows)))
        except StopIteration:
            cols = {}
        if "unit_id" in cols and "session_id" in cols:
            for row in rows:
                out.stats.practice_rows += 1
                uid = (str(row[cols["unit_id"]]).strip()
                       if cols["unit_id"] < len(row) and row[cols["unit_id"]] else "")
                sid = (str(row[cols["session_id"]]).strip()
                       if cols["session_id"] < len(row) and row[cols["session_id"]] else "")
                no = by_session_id.get(sid)
                if uid and no:
                    out.session_of_unit[uid] = no
                else:
                    out.stats.unmapped_practice_rows += 1
            out.stats.units_mapped = len(out.session_of_unit)
        else:
            out.stats.errors.append(
                f"{PRACTICE_SHEET}: needs both 'Unit ID' and 'Session ID' columns")
    else:
        out.stats.errors.append(f"no '{PRACTICE_SHEET}' sheet in {Path(path).name}")

    return out


def outline_records(outline: Outline, course: str,
                    source_file: str = "") -> list[ContentRecord]:
    """The slide-level text as ContentRecords, so the extractor can read it.

    Emitted as `markdown` evidence because that is what it is - authored prose - and it
    is what `InventoryBuilder.feed_prose` reads. The `field_path` names the workbook,
    sheet, row and column instead of a JSON path, which is what tells the detail panel
    to explain that there is no export excerpt to show rather than rendering blank.

    A tool named here is taught on a slide. That is weaker evidence than an import in a
    reference solution and stronger than a passing mention in prose, and it is the ONLY
    evidence for tools that never appear in the export at all.
    """
    out: list[ContentRecord] = []
    for s in outline.sessions:
        for col, text in (("Outline", s.outline), ("Key Takeaways", s.key_takeaways)):
            if len(text.strip()) < 3:
                continue
            out.append(ContentRecord(
                course=course, topic_name=s.topic_name, unit_id=s.session_id,
                unit_name=s.session_name or f"Session {s.session_no}",
                unit_type="SESSION", content_id="", object_type="SESSION_PPT",
                content_type="TEXT", title=col, body_text=text,
                field_path=f"{outline.stats.workbook}::{OUTLINE_SHEET}::{col}::row{s.row}",
                evidence_source="markdown", source_file=source_file or "workbook",
                session_no=s.session_no))
    return out
